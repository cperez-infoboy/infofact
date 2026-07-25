"""End-to-end test against the real RFP xlsx: ingest -> extract (real LLM) ->
verify_spans -> recall vs the requirement Descriptions in the workbook (ground
truth). Measures whether the pipeline recovers the client's stated requirements.

Assumption (documented): ground truth = the Descripción column (col D) of both
NECESIDADES sheets. A requirement is "recovered" if some extracted statement has
embedding cosine >= THRESH against a GT description.
"""
import asyncio
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from backend.agents.pipelines.consolidation import embed_texts
from backend.agents.pipelines.extraction import extract_all
from backend.agents.pipelines.ingestion import ingest_document

PATH = Path("docs/testing/1. Requerimientos tecnicos- funcionales.xlsx")
THRESH = 0.70
PROJECT_NAME = "TCC - Gestion de rutas y entregas"
PROJECT_DESC = (
    "Aplicacion para gestion de rutas (recogidas y entregas), cotizacion de "
    "servicios, recaudo, reportes BI y modulo administrativo. Necesidades "
    "funcionales y tecnicas para licitacion."
)


def load_ground_truth() -> list[str]:
    wb = openpyxl.load_workbook(PATH, data_only=False)
    descs: list[str] = []
    for sheet in ("NECESIDADES FUNCIONALES", "NECESIDADES TÉCNICAS"):
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=71, max_row=ws.max_row, values_only=True):
            desc = r[3] if len(r) > 3 else None
            if desc and str(desc).strip():
                descs.append(str(desc).strip())
    return descs


async def main():
    print("ingesting xlsx ...")
    chunks, smap = ingest_document(PATH)
    print(f"  chunks={len(chunks)}  full_text={len(smap.full_text)} chars")

    print("extracting requirements (real LLM, concurrency=4) ...")
    doc_texts = {str(PATH): smap.full_text}
    items = await extract_all(
        chunks,
        doc_texts=doc_texts,
        project_name=PROJECT_NAME,
        project_description=PROJECT_DESC,
        concurrency=4,
    )
    verified = sum(1 for i in items if i.span_verified)
    print(f"  extracted={len(items)}  span_verified={verified}")

    print("loading ground truth + measuring recall ...")
    gt = load_ground_truth()
    print(f"  GT descriptions: {len(gt)}")

    gt_vecs = embed_texts(gt)
    ex_vecs = embed_texts([i.statement for i in items]) if items else None
    if ex_vecs is None:
        print("  NO items extracted; recall=0")
        return
    sim = ex_vecs @ gt_vecs.T               # rows=extracted, cols=GT
    best_per_gt = sim.max(axis=0)           # best extracted match per GT
    matched = int((best_per_gt >= THRESH).sum())
    recall = matched / len(gt) if gt else 0.0
    best_per_ex = sim.max(axis=1)           # best GT match per extracted
    ex_with_gt = int((best_per_ex >= THRESH).sum())
    precision = ex_with_gt / len(items) if items else 0.0

    print()
    print(f"=== RESULTS ===")
    print(f"  extracted items:        {len(items)}")
    print(f"  span_verified:          {verified}/{len(items)} ({verified/max(1,len(items)):.0%})")
    print(f"  GT descriptions:        {len(gt)}")
    print(f"  recall@{THRESH}:          {matched}/{len(gt)} = {recall:.1%}")
    print(f"  precision@{THRESH}:       {ex_with_gt}/{len(items)} = {precision:.1%}")

    # show a few GT that were NOT matched (misses)
    misses = [gt[j] for j in range(len(gt)) if best_per_gt[j] < THRESH]
    print(f"  --- {len(misses)} GT misses (first 5) ---")
    for m in misses[:5]:
        print(f"    MISS: {m[:100]!r}")

    # show extracted items with no GT match (potential hallucination / noise)
    noise = [(i.statement, float(best_per_ex[k])) for k, i in enumerate(items) if best_per_ex[k] < THRESH]
    print(f"  --- {len(noise)} extracted without GT match (first 5) ---")
    for stmt, score in noise[:5]:
        print(f"    NOISE(best={score:.2f}): {stmt[:90]!r}")

    # persist results for human review
    _dump(items, best_per_ex, THRESH, gt, PATH)


def _dump(items, best_per_ex, thresh, gt, source_path: Path) -> None:
    import json
    out_dir = Path("docs/testing")
    json_path = out_dir / "extraccion_resultado.json"
    md_path = out_dir / "extraccion_resultado.md"

    # raw JSON
    json_path.write_text(
        json.dumps([i.model_dump() for i in items], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    # human-readable markdown, grouped by section
    verified = sum(1 for i in items if i.span_verified)
    by_section: dict[str, list] = {}
    for k, it in enumerate(items):
        by_section.setdefault(it.section or "(sin sección)", []).append((k, it))

    lines = [
        f"# Extracción de requerimientos — {source_path.name}",
        "",
        f"- Extraídos: **{len(items)}**",
        f"- Span verificados: **{verified}/{len(items)}** ({verified/max(1,len(items)):.0%})",
        f"- GT (col Descripción): **{len(gt)}**  ·  recall@{thresh}: ver consola",
        "",
        "Leyenda: ✓ span verificado · ⚠ span NO verificado (posible alucinación) · "
        "Noise = sin match contra GT (ruido de plantilla, no alucinación).",
        "",
    ]
    for section in sorted(by_section):
        entries = by_section[section]
        lines.append(f"## {section}  ({len(entries)})")
        lines.append("")
        for n, (k, it) in enumerate(entries, 1):
            badge = "✓" if it.span_verified else "⚠"
            noise_tag = " · NOISE" if best_per_ex[k] < thresh else ""
            page = f" p.{it.page}" if it.page else ""
            lines.append(f"### {badge} {it.id or n}  · conf {it.confidence:.2f}{page}{noise_tag}")
            lines.append(f"**Statement:** {it.statement}")
            lines.append("")
            lines.append(f"> **Source:** {it.source_span}")
            lines.append("")
    md_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"\n  dump: {md_path}  +  {json_path}")


if __name__ == "__main__":
    asyncio.run(main())
